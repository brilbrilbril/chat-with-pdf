import openpyxl

from src.domain.interfaces.extractor import IBaseExtractor
from src.domain.entities.chunk import RawChunk


class XLSXExtractor(IBaseExtractor):
    def __init__(self, rows_per_chunk: int = 5):
        self.rows_per_chunk = rows_per_chunk

    def extract(self, file_path: str) -> list[RawChunk]:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        raw_chunks: list[RawChunk] = []
        chunk_index = 0

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            all_rows = list(ws.iter_rows(values_only=True))

            if not all_rows:
                continue

            header = all_rows[0]
            header_text = " | ".join(str(c) if c is not None else "" for c in header)
            data_rows = all_rows[1:]

            if not data_rows:
                raw_chunks.append(RawChunk(
                    text=f"Sheet: {sheet_name}\n{header_text}",
                    sheet_name=sheet_name,
                    row_start=1,
                    row_end=1,
                    chunk_index=chunk_index,
                ))
                chunk_index += 1
                continue

            for batch_start in range(0, len(data_rows), self.rows_per_chunk):
                batch = data_rows[batch_start: batch_start + self.rows_per_chunk]
                row_start = batch_start + 2
                row_end = row_start + len(batch) - 1

                row_texts = []
                for row in batch:
                    row_text = " | ".join(str(c) if c is not None else "" for c in row)
                    row_texts.append(row_text)

                chunk_text = (
                    f"Sheet: {sheet_name} | Rows {row_start}-{row_end}\n"
                    f"{header_text}\n"
                    + "\n".join(row_texts)
                )

                raw_chunks.append(RawChunk(
                    text=chunk_text,
                    sheet_name=sheet_name,
                    row_start=row_start,
                    row_end=row_end,
                    chunk_index=chunk_index,
                ))
                chunk_index += 1

        wb.close()
        return raw_chunks