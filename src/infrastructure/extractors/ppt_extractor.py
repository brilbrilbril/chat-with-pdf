from pptx import Presentation

from src.domain.interfaces.extractor import IBaseExtractor
from src.domain.entities.chunk import RawChunk
from src.infrastructure.extractors.chunker import recursive_split


class PPTXExtractor(IBaseExtractor):
    def __init__(self, chunk_size: int = 500, chunk_overlap: int = 50):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap

    def extract(self, file_path: str) -> list[RawChunk]:
        prs = Presentation(file_path)
        raw_chunks: list[RawChunk] = []
        chunk_index = 0

        for slide_num, slide in enumerate(prs.slides, start=1):
            parts: list[str] = []

            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            parts.append(text)

                if shape.has_table:
                    rows = []
                    for row in shape.table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        rows.append(" | ".join(cells))
                    if rows:
                        parts.append("\n".join(rows))

            slide_text = "\n\n".join(parts).strip()
            if not slide_text:
                continue

            splits = recursive_split(slide_text, self.chunk_size, self.chunk_overlap)
            for split in splits:
                raw_chunks.append(RawChunk(
                    text=split,
                    slide_number=slide_num,
                    chunk_index=chunk_index,
                ))
                chunk_index += 1

        return raw_chunks